from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_DIR = Path("/Users/rishi/Documents/Test Case Generator/docs")
PRD_PATH = OUTPUT_DIR / "test-case-generator-prd.docx"
TECH_SPEC_PATH = OUTPUT_DIR / "test-case-generator-technical-specification.docx"

ACCENT = RGBColor(46, 116, 181)
ACCENT_DARK = RGBColor(31, 77, 120)
TEXT = RGBColor(34, 34, 34)
MUTED = RGBColor(90, 90, 90)
TABLE_FILL = "F2F4F7"
TABLE_RULE = "DADCE0"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
      tc_w = OxmlElement("w:tcW")
      tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(int(inches * 1440)))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), TABLE_RULE)


def apply_run_font(run, size, bold=False, color=TEXT, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def build_base_document(title, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")

    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.color.rgb = ACCENT
    styles["Heading 1"].paragraph_format.space_before = Pt(16)
    styles["Heading 1"].paragraph_format.space_after = Pt(8)

    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 2"].font.color.rgb = ACCENT
    styles["Heading 2"].paragraph_format.space_before = Pt(12)
    styles["Heading 2"].paragraph_format.space_after = Pt(6)

    styles["Heading 3"].font.size = Pt(12)
    styles["Heading 3"].font.color.rgb = ACCENT_DARK
    styles["Heading 3"].paragraph_format.space_before = Pt(8)
    styles["Heading 3"].paragraph_format.space_after = Pt(4)

    if "Callout" not in styles:
        callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
        callout.base_style = normal
        callout.font.name = "Calibri"
        callout.font.size = Pt(10.5)
        callout.font.color.rgb = TEXT
        callout.paragraph_format.space_before = Pt(0)
        callout.paragraph_format.space_after = Pt(8)
        callout.paragraph_format.left_indent = Inches(0.15)
        callout.paragraph_format.right_indent = Inches(0.15)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.add_run("Test Case Generator Platform")
    apply_run_font(header_run, 9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("Internal Working Draft")
    apply_run_font(footer_run, 9, color=MUTED)

    title_paragraph = doc.add_paragraph()
    title_paragraph.paragraph_format.space_before = Pt(0)
    title_paragraph.paragraph_format.space_after = Pt(4)
    title_run = title_paragraph.add_run(title)
    apply_run_font(title_run, 22, bold=True, color=TEXT)

    subtitle_paragraph = doc.add_paragraph()
    subtitle_paragraph.paragraph_format.space_before = Pt(0)
    subtitle_paragraph.paragraph_format.space_after = Pt(14)
    subtitle_run = subtitle_paragraph.add_run(subtitle)
    apply_run_font(subtitle_run, 12, color=MUTED)

    return doc


def add_body_paragraph(doc, text):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(text)
    apply_run_font(run, 11)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.15
        run = paragraph.add_run(item)
        apply_run_font(run, 11)


def add_numbered_steps(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.15
        run = paragraph.add_run(item)
        apply_run_font(run, 11)


def add_key_value_table(doc, rows, widths=(1.6, 4.9)):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    set_table_borders(table)
    for label, value in rows:
        row = table.add_row()
        left, right = row.cells
        set_cell_width(left, widths[0])
        set_cell_width(right, widths[1])
        left.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        right.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(left, TABLE_FILL)
        p1 = left.paragraphs[0]
        p2 = right.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(label)
        r2 = p2.add_run(value)
        apply_run_font(r1, 10.5, bold=True)
        apply_run_font(r2, 10.5)
    doc.add_paragraph()


def add_matrix_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    set_table_borders(table)
    header_row = table.rows[0]
    set_repeat_table_header(header_row)

    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        set_cell_width(cell, widths[index])
        set_cell_shading(cell, TABLE_FILL)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        apply_run_font(run, 10, bold=True)

    for row_values in rows:
        row = table.add_row()
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            set_cell_width(cell, widths[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            apply_run_font(run, 9.5)

    doc.add_paragraph()


def build_prd():
    doc = build_base_document(
        "Product Requirements Document",
        "AI-powered test case generator for PRD-led QA planning",
    )

    add_key_value_table(
        doc,
        [
            ("Document owner", "Product Management"),
            ("Primary stakeholders", "QA, Engineering, Product, Release Management"),
            ("Target release", "Phase 1 production rollout"),
            ("Product scope", "Upload PRD and knowledge-base documents, analyze them with Gemini, and generate QA-ready test cases"),
        ],
    )

    doc.add_heading("1. Business context", level=1)
    add_body_paragraph(
        doc,
        "QA teams often receive feature PRDs late in the lifecycle and must manually translate them into coverage plans under tight time pressure. The current process is document-heavy, inconsistent across reviewers, and dependent on individual domain memory to identify validation, regression, permission, integration, and observability scenarios."
    )
    add_body_paragraph(
        doc,
        "The Test Case Generator product addresses this by letting teams upload a source PRD and relevant knowledge-base material, then using a real LLM workflow to extract intent, reconcile supporting context, and propose high-signal QA coverage. The goal is not to replace human QA judgment, but to dramatically shorten first-draft creation and reduce missed coverage."
    )
    add_bullets(
        doc,
        [
            "Primary users are QA engineers, SDETs, product analysts, and engineering managers reviewing feature readiness.",
            "Source artifacts include PRDs, requirement notes, process documents, configuration guides, policy documents, and integration notes.",
            "The initial release focuses on feature-specific testcase drafting rather than full test management, execution tracking, or defect lifecycle workflows.",
        ],
    )

    doc.add_heading("2. Benefits", level=1)
    add_bullets(
        doc,
        [
            "Reduces the time required to prepare a first-draft QA coverage plan from hours to minutes.",
            "Improves consistency in how PRDs are interpreted across QA reviewers and teams.",
            "Surfaces risky areas such as missing edge cases, implied dependencies, and likely regressions earlier in the release cycle.",
            "Preserves team knowledge by using uploaded supporting documents rather than relying purely on individual memory.",
            "Produces structured output that can later be exported into test management or workflow tooling.",
        ],
    )

    doc.add_heading("3. Concept Note", level=1)
    add_body_paragraph(
        doc,
        "The workflow begins with a user uploading a single PRD and any number of contextual knowledge-base documents. The backend parses PDF and DOCX files server-side, normalizes the extracted text, and submits the combined context to Gemini using a structured JSON schema."
    )
    add_numbered_steps(
        doc,
        [
            "User lands on the QA workspace and uploads the PRD plus optional supporting documents.",
            "Frontend sends all files to the backend through a multipart request.",
            "Backend extracts readable text from PDF, DOCX, and text-like formats and trims content to safe model limits.",
            "Gemini analyzes the PRD as the primary source, uses the supporting documents as secondary context, and returns a structured feature summary, assumptions, risks, and prioritized test cases.",
            "Frontend renders the resulting suite with criticality badges and expandable orchestration steps for each test case.",
            "User reviews the output and can iterate by replacing the PRD, adding more knowledge base context, or removing noisy supporting files.",
        ],
    )
    add_body_paragraph(
        doc,
        "A production-grade roadmap would add persistence, run history, export options, and collaboration features, but the MVP centers on high-quality ingestion, trustworthy generation, and a review-friendly presentation layer."
    )

    doc.add_heading("4. Table details", level=1)
    add_body_paragraph(
        doc,
        "The current prototype is stateless, but the production implementation should persist uploads, parsing metadata, generation runs, and generated artifacts. The following relational tables are proposed for Phase 2 persistence."
    )
    add_matrix_table(
        doc,
        ["Table", "Purpose", "Primary relationships"],
        [
            ("features", "Represents a PRD-driven feature analysis request or feature workspace.", "Parent record for documents and generation runs"),
            ("documents", "Stores uploaded files and extraction metadata for PRDs and knowledge-base files.", "Belongs to features"),
            ("generation_runs", "Tracks each analysis run against a feature workspace.", "Belongs to features"),
            ("generated_test_cases", "Stores testcase-level output for a generation run.", "Belongs to generation_runs"),
            ("generated_test_steps", "Stores ordered execution steps for each generated test case.", "Belongs to generated_test_cases"),
        ],
        [1.1, 2.0, 3.4],
    )
    add_matrix_table(
        doc,
        ["Table", "Field", "Field type", "Description"],
        [
            ("features", "id", "UUID", "Primary key for the feature workspace."),
            ("features", "title", "VARCHAR(255)", "Human-readable feature or PRD title."),
            ("features", "status", "VARCHAR(32)", "Current lifecycle state such as draft, analyzed, or archived."),
            ("features", "created_at", "TIMESTAMP", "Record creation timestamp."),
            ("features", "updated_at", "TIMESTAMP", "Last modified timestamp."),
            ("documents", "id", "UUID", "Primary key for the uploaded document."),
            ("documents", "feature_id", "UUID", "Foreign key linking the document to a feature workspace."),
            ("documents", "document_type", "VARCHAR(32)", "PRD or KNOWLEDGE_BASE."),
            ("documents", "original_filename", "VARCHAR(255)", "Name of the uploaded file."),
            ("documents", "mime_type", "VARCHAR(128)", "Uploaded content type."),
            ("documents", "storage_path", "TEXT", "Pointer to object storage or managed file store."),
            ("documents", "extracted_text", "LONGTEXT", "Normalized text extracted by the parser."),
            ("documents", "excerpt", "TEXT", "Short preview for UI display."),
            ("documents", "parse_status", "VARCHAR(32)", "PARSED, FAILED, or PARTIAL."),
            ("documents", "created_at", "TIMESTAMP", "Record creation timestamp."),
            ("generation_runs", "id", "UUID", "Primary key for each generation attempt."),
            ("generation_runs", "feature_id", "UUID", "Foreign key to the feature workspace."),
            ("generation_runs", "model_name", "VARCHAR(128)", "LLM model used for the run."),
            ("generation_runs", "prompt_version", "VARCHAR(64)", "Prompt template version used for traceability."),
            ("generation_runs", "status", "VARCHAR(32)", "QUEUED, RUNNING, COMPLETED, or FAILED."),
            ("generation_runs", "summary", "TEXT", "Generated short feature summary."),
            ("generation_runs", "risks_json", "JSON", "Array of highlighted risks."),
            ("generation_runs", "assumptions_json", "JSON", "Array of model assumptions."),
            ("generation_runs", "source_highlights_json", "JSON", "Array of source references shown in UI."),
            ("generation_runs", "created_at", "TIMESTAMP", "Record creation timestamp."),
            ("generated_test_cases", "id", "UUID", "Primary key for a generated test case."),
            ("generated_test_cases", "generation_run_id", "UUID", "Foreign key to the generation run."),
            ("generated_test_cases", "title", "VARCHAR(255)", "Short test case title."),
            ("generated_test_cases", "criticality", "VARCHAR(16)", "Low, Medium, or High."),
            ("generated_test_cases", "rationale", "TEXT", "Reason the case matters."),
            ("generated_test_cases", "coverage_area", "VARCHAR(128)", "Logical category such as validation or regression."),
            ("generated_test_steps", "id", "UUID", "Primary key for a generated step."),
            ("generated_test_steps", "generated_test_case_id", "UUID", "Foreign key to a generated test case."),
            ("generated_test_steps", "step_order", "INTEGER", "Display order for the step."),
            ("generated_test_steps", "instruction_text", "TEXT", "Actionable instruction for the tester."),
        ],
        [1.1, 1.4, 1.1, 2.9],
    )

    doc.add_heading("5. User experience and flows", level=1)
    doc.add_heading("Screen 1: QA Workspace Landing", level=2)
    add_body_paragraph(
        doc,
        "The landing screen is the main working canvas. It combines upload actions, a concise explanation of the generation flow, and a contextual summary panel that shows what has already been attached."
    )
    add_bullets(
        doc,
        [
            "Inputs available: upload PRD, upload knowledge-base documents, remove attached knowledge-base files, and trigger generation.",
            "Visible information: workspace title, expected output summary, current PRD filename, count of supporting documents, and any in-browser text preview available for text-like formats.",
            "Validation: generation button remains disabled until a PRD is attached.",
        ],
    )

    doc.add_heading("Screen 2: PRD Upload Flow", level=2)
    add_bullets(
        doc,
        [
            "Supported input formats: PDF, DOCX, TXT, MD, and RTF for PRDs.",
            "The system accepts one PRD per run and replaces the current PRD if a new file is uploaded.",
            "The UI confirms successful upload with filename-level feedback and resets old generated results to avoid stale analysis.",
        ],
    )

    doc.add_heading("Screen 3: Knowledge Base Upload Flow", level=2)
    add_bullets(
        doc,
        [
            "Supported input formats: PDF, DOCX, TXT, MD, CSV, and JSON.",
            "Users may upload multiple documents in one action and continue adding more in later actions.",
            "Each uploaded knowledge-base file is shown with filename and a short preview if the browser can read it directly.",
            "Users can remove individual supporting documents before regeneration.",
        ],
    )

    doc.add_heading("Screen 4: Analysis in Progress", level=2)
    add_bullets(
        doc,
        [
            "The main panel switches to a loading state with progress messaging.",
            "The user is informed that the system is parsing documents, identifying risks, and drafting test cases.",
            "The sidebar continues showing the PRD and knowledge-base attachment status for orientation.",
        ],
    )

    doc.add_heading("Screen 5: Generated Test Case Review", level=2)
    add_bullets(
        doc,
        [
            "A short feature summary is displayed first, grounded in PRD understanding and supporting context.",
            "Risk scan, assumptions, and source context are visible so users can judge output trustworthiness.",
            "Each generated test case shows title, rationale, coverage area, and criticality badge.",
            "Each test case can be expanded to reveal step-by-step execution guidance.",
            "Expected user action: review, refine source inputs, and rerun generation as needed.",
        ],
    )

    doc.add_heading("Error and edge-state expectations", level=2)
    add_bullets(
        doc,
        [
            "If the PRD is missing, the system must block generation and show a clear corrective message.",
            "If document parsing fails, the user should see which file failed and why.",
            "If the model returns an error or empty output, the frontend must preserve the uploaded context and allow retry without forcing re-upload.",
        ],
    )

    doc.save(PRD_PATH)


def build_tech_spec():
    doc = build_base_document(
        "Technical Specification Document",
        "Architecture and API design for the Test Case Generator backend",
    )

    add_key_value_table(
        doc,
        [
            ("System owner", "Engineering"),
            ("Runtime stack", "TypeScript, Express, Vite, Gemini API"),
            ("Document ingestion", "Multer, pdf-parse, mammoth"),
            ("Delivery model", "Single repo with frontend and backend running in parallel"),
        ],
    )

    doc.add_heading("1. Executive Summary", level=1)
    add_body_paragraph(
        doc,
        "The Test Case Generator backend provides a document-ingestion and inference service that accepts one PRD plus multiple contextual knowledge-base files, extracts the underlying text, sends the combined context to Gemini, and returns structured QA output to the frontend."
    )
    add_body_paragraph(
        doc,
        "The MVP is intentionally lightweight and stateless. It is implemented as an Express API in the same repository as the Vite frontend, with shared TypeScript types used to keep the request-response contract stable. The system focuses on three engineering responsibilities: reliable file upload handling, deterministic document parsing, and trustworthy structured LLM output validation."
    )

    doc.add_heading("2. System Architechture", level=1)
    add_body_paragraph(
        doc,
        "The system follows a browser-client to API-server to LLM-provider flow. The frontend handles user interaction and file selection, the backend handles document extraction and orchestration, and Gemini performs structured reasoning over the assembled context."
    )
    add_numbered_steps(
        doc,
        [
            "User selects a PRD and optional knowledge-base files in the React frontend.",
            "Frontend constructs a multipart form request and submits it to POST /api/generate-test-cases.",
            "Express receives the request and Multer stores uploaded files in memory for immediate processing.",
            "Document parsing service detects file type and extracts text from PDF, DOCX, or text-like files.",
            "The generation service trims the extracted content to safe token-oriented limits and builds a prompt with the PRD as the primary context source.",
            "Gemini returns JSON constrained to the expected schema for summaries, risks, assumptions, test cases, and test steps.",
            "Backend validates the JSON with Zod and returns the resulting payload to the frontend.",
            "Frontend renders the structured output in the testcase explorer.",
        ],
    )
    add_matrix_table(
        doc,
        ["Layer", "Primary responsibility", "Implementation notes"],
        [
            ("Frontend", "User interaction and display", "React + Vite UI, uploads files, renders suite, handles loading and error states"),
            ("API layer", "Transport and validation", "Express routes, CORS, JSON responses, health endpoint, multipart handling"),
            ("Parsing layer", "Document text extraction", "pdf-parse for PDF, mammoth for DOCX, direct text extraction for TXT/MD/CSV/JSON/RTF"),
            ("Inference layer", "LLM orchestration", "Gemini prompt assembly, structured output request, schema-constrained generation"),
            ("Shared contract", "Type safety across client and server", "Shared GeneratedSuite types used on both frontend and backend"),
        ],
        [1.0, 2.0, 3.5],
    )

    doc.add_heading("3. Data Model", level=1)
    add_body_paragraph(
        doc,
        "The live MVP does not yet persist runs, but the following data model should be used when the system graduates from transient processing to durable storage. The model separates uploaded source context, generation attempts, and the generated artifacts themselves."
    )
    add_matrix_table(
        doc,
        ["Table", "Description"],
        [
            ("features", "Root feature workspace created from an uploaded PRD and optional supporting knowledge base."),
            ("documents", "Stores metadata, extracted text, and storage pointers for PRD and supporting documents."),
            ("generation_runs", "Tracks each attempt to generate a test suite, including model and status metadata."),
            ("generated_test_cases", "Persists testcase-level results for a given generation run."),
            ("generated_test_steps", "Stores ordered steps associated with one generated test case."),
        ],
        [1.4, 5.1],
    )
    add_bullets(
        doc,
        [
            "A feature can have many documents and many generation runs.",
            "A generation run belongs to one feature and produces many test cases.",
            "A generated test case belongs to one generation run and has many ordered test steps.",
            "Documents should be tagged by type so the PRD remains distinguishable from supporting knowledge-base content during audits and reruns.",
        ],
    )
    add_matrix_table(
        doc,
        ["Entity", "Key fields", "Notes"],
        [
            ("features", "id, title, status, created_at, updated_at", "Used to group all analysis activity for one feature scope."),
            ("documents", "id, feature_id, document_type, original_filename, mime_type, extracted_text, parse_status", "Stores both upload metadata and parser results."),
            ("generation_runs", "id, feature_id, model_name, prompt_version, status, summary, risks_json, assumptions_json", "Supports observability and re-run traceability."),
            ("generated_test_cases", "id, generation_run_id, title, criticality, rationale, coverage_area", "One row per surfaced QA scenario."),
            ("generated_test_steps", "id, generated_test_case_id, step_order, instruction_text", "One row per expandable execution step."),
        ],
        [1.2, 2.4, 2.9],
    )

    doc.add_heading("4. Backend API specifications (if any)", level=1)
    add_heading = doc.add_heading
    add_heading("GET /api/health", level=2)
    add_bullets(
        doc,
        [
            "Purpose: lightweight uptime and readiness check for the API service.",
            "Request body: none.",
            'Success response: `200 OK` with payload `{ "status": "ok" }`.',
        ],
    )

    add_heading("POST /api/generate-test-cases", level=2)
    add_body_paragraph(
        doc,
        "This is the primary application endpoint. It accepts uploaded source documents, performs extraction and LLM orchestration, and returns a fully structured testcase suite."
    )
    add_matrix_table(
        doc,
        ["Aspect", "Specification"],
        [
            ("Method", "POST"),
            ("Path", "/api/generate-test-cases"),
            ("Content type", "multipart/form-data"),
            ("Required file field", "prd"),
            ("Optional repeated file field", "knowledgeBase"),
            ("Maximum upload strategy", "Memory-backed upload with file count and file size limits enforced in Multer"),
            ("Supported parsing types", "PDF, DOCX, TXT, MD, CSV, JSON, RTF"),
        ],
        [1.8, 4.7],
    )
    add_heading("Request contract", level=3)
    add_bullets(
        doc,
        [
            "Exactly one file must be sent in the `prd` field.",
            "Zero or more files may be sent in the `knowledgeBase` field.",
            "Files are parsed server-side, not in the browser.",
        ],
    )

    add_heading("Success response contract", level=3)
    add_matrix_table(
        doc,
        ["Field", "Type", "Description"],
        [
            ("featureName", "string", "Generated human-readable feature title."),
            ("shortSummary", "string", "Concise description of the feature based on the PRD."),
            ("risks", "string[]", "List of noteworthy QA risks or potential blind spots."),
            ("assumptions", "string[]", "Model assumptions or inferred constraints."),
            ("sourceHighlights", "string[]", "Key references surfaced from source documents."),
            ("cases", "GeneratedTestCase[]", "Array of test cases with criticality, rationale, coverage area, and steps."),
        ],
        [1.4, 1.4, 3.7],
    )

    add_heading("GeneratedTestCase schema", level=3)
    add_matrix_table(
        doc,
        ["Field", "Type", "Description"],
        [
            ("id", "string", "Stable identifier for the testcase within the response."),
            ("title", "string", "Short testcase name."),
            ("criticality", "Low | Medium | High", "Priority or impact marker for QA review."),
            ("rationale", "string", "Why the testcase matters."),
            ("coverageArea", "string", "Logical grouping such as validation or regression."),
            ("steps", "TestStep[]", "Ordered steps used by the expandable UI."),
        ],
        [1.5, 1.7, 3.3],
    )

    add_heading("Failure modes", level=3)
    add_bullets(
        doc,
        [
            "400 Bad Request when the PRD is missing or the file type is unsupported.",
            "500 Internal Server Error when parsing fails unexpectedly, Gemini returns unusable output, or the API key is missing.",
            "Error payload shape: `{ \"error\": \"human-readable message\" }`.",
        ],
    )

    doc.save(TECH_SPEC_PATH)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    build_prd()
    build_tech_spec()
    print(PRD_PATH)
    print(TECH_SPEC_PATH)


if __name__ == "__main__":
    main()
